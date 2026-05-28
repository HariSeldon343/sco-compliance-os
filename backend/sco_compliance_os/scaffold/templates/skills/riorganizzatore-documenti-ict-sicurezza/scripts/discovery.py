#!/usr/bin/env python3
"""Phase 1: Discovery - Scan folder recursively, extract metadata and text previews."""

import argparse
import hashlib
import json
import mimetypes
import os
import sys
from datetime import datetime
from pathlib import Path

DEFAULT_EXCLUDE = [".git", "node_modules", "__pycache__", ".DS_Store", ".Trash", "Thumbs.db"]
TEXT_EXTENSIONS = {".txt", ".md", ".rst", ".html", ".htm", ".css", ".js", ".py", ".sh", ".json", ".xml", ".yaml", ".yml", ".csv", ".tsv", ".log", ".ini", ".cfg", ".conf"}
MAX_PREVIEW = 2000


def sha256_file(path, chunk_size=65536):
    h = hashlib.sha256()
    try:
        with open(path, "rb") as f:
            while True:
                chunk = f.read(chunk_size)
                if not chunk:
                    break
                h.update(chunk)
        return h.hexdigest()
    except (PermissionError, OSError):
        return None


def extract_text_docx(path):
    try:
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        title = doc.core_properties.title or ""
        author = doc.core_properties.author or ""
        return text[:MAX_PREVIEW], title, author
    except Exception as e:
        return f"[extraction error: {e}]", "", ""


def extract_text_pdf(path):
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages):
                if len("\n".join(text_parts)) >= MAX_PREVIEW:
                    break
                t = page.extract_text()
                if t:
                    text_parts.append(t)
        text = "\n".join(text_parts)
        return text[:MAX_PREVIEW], "", ""
    except Exception:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(path)
            text_parts = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
                if len("\n".join(text_parts)) >= MAX_PREVIEW:
                    break
            info = reader.metadata
            title = (info.title or "") if info else ""
            author = (info.author or "") if info else ""
            return "\n".join(text_parts)[:MAX_PREVIEW], title, author
        except Exception as e:
            return f"[extraction error: {e}]", "", ""


def extract_text_xlsx(path):
    try:
        import openpyxl  # type: ignore[import-untyped]
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        text_parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            text_parts.append(f"[Sheet: {sheet_name}]")
            for row in ws.iter_rows(max_row=50, values_only=True):
                vals = [str(c) for c in row if c is not None]
                if vals:
                    text_parts.append(" | ".join(vals))
                if len("\n".join(text_parts)) >= MAX_PREVIEW:
                    break
        wb.close()
        return "\n".join(text_parts)[:MAX_PREVIEW], "", ""
    except Exception as e:
        return f"[extraction error: {e}]", "", ""


def extract_text_pptx(path):
    try:
        from pptx import Presentation
        prs = Presentation(path)
        text_parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text_parts.append(para.text)
            if len("\n".join(text_parts)) >= MAX_PREVIEW:
                break
        return "\n".join(text_parts)[:MAX_PREVIEW], "", ""
    except Exception as e:
        return f"[extraction error: {e}]", "", ""


def extract_text_csv(path):
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read(MAX_PREVIEW)
        return text, "", ""
    except Exception as e:
        return f"[extraction error: {e}]", "", ""


def extract_text_plain(path):
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read(MAX_PREVIEW)
        return text, "", ""
    except Exception as e:
        return f"[extraction error: {e}]", "", ""


def extract_text(path):
    ext = Path(path).suffix.lower()
    if ext in (".docx",):
        return extract_text_docx(path)
    elif ext in (".pdf",):
        return extract_text_pdf(path)
    elif ext in (".xlsx", ".xls"):
        return extract_text_xlsx(path)
    elif ext in (".pptx",):
        return extract_text_pptx(path)
    elif ext in (".csv", ".tsv"):
        return extract_text_csv(path)
    elif ext in TEXT_EXTENSIONS:
        return extract_text_plain(path)
    elif ext in (".odt", ".ods", ".odp", ".doc", ".ppt"):
        return f"[format {ext} - use pandoc or LibreOffice to convert]", "", ""
    else:
        return "", "", ""


def scan_directory(source_path, exclude_patterns=None, max_depth=0):
    if exclude_patterns is None:
        exclude_patterns = DEFAULT_EXCLUDE
    source = Path(source_path).resolve()
    if not source.is_dir():
        print(f"Error: {source} is not a directory", file=sys.stderr)
        sys.exit(1)

    inventory = []
    file_count = 0
    skipped = 0

    for root, dirs, files in os.walk(source, followlinks=False):
        rel_root = Path(root).relative_to(source)
        depth = len(rel_root.parts)
        if max_depth > 0 and depth > max_depth:
            dirs.clear()
            continue

        dirs[:] = [d for d in dirs if d not in exclude_patterns and not d.startswith(".")]

        for fname in files:
            if fname in exclude_patterns or fname.startswith("."):
                skipped += 1
                continue

            fpath = Path(root) / fname
            try:
                stat = fpath.stat()
            except (PermissionError, OSError):
                skipped += 1
                continue

            file_count += 1
            rel_path = str(fpath.relative_to(source))
            ext = fpath.suffix.lower()
            mime, _ = mimetypes.guess_type(str(fpath))

            print(f"  [{file_count}] Scanning: {rel_path}", file=sys.stderr)

            checksum = sha256_file(fpath)
            text_preview, title, author = extract_text(fpath)

            entry = {
                "id": file_count,
                "file_path": rel_path,
                "file_name": fname,
                "file_ext": ext,
                "file_size": stat.st_size,
                "created_at": datetime.fromtimestamp(stat.st_ctime).isoformat(),
                "modified_at": datetime.fromtimestamp(stat.st_mtime).isoformat(),
                "mime_type": mime or "application/octet-stream",
                "checksum_sha256": checksum,
                "title": title,
                "author": author,
                "text_preview": text_preview,
                "folder_path": str(rel_root) if str(rel_root) != "." else ""
            }
            inventory.append(entry)

    return {
        "source_path": str(source),
        "scan_date": datetime.now().isoformat(),
        "total_files": file_count,
        "skipped_files": skipped,
        "files": inventory
    }


def main():
    parser = argparse.ArgumentParser(description="Phase 1: Discovery - Scan and inventory documents")
    parser.add_argument("source_path", help="Path to the source folder to scan")
    parser.add_argument("--output", "-o", default="inventory.json", help="Output JSON file (default: inventory.json)")
    parser.add_argument("--exclude", default=",".join(DEFAULT_EXCLUDE), help="Comma-separated exclude patterns")
    parser.add_argument("--max-depth", type=int, default=0, help="Max recursion depth (0=unlimited)")
    args = parser.parse_args()

    exclude = [x.strip() for x in args.exclude.split(",")]
    print(f"Scanning: {args.source_path}", file=sys.stderr)
    result = scan_directory(args.source_path, exclude, args.max_depth)

    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\nDiscovery complete:", file=sys.stderr)
    print(f"  Total files: {result['total_files']}", file=sys.stderr)
    print(f"  Skipped: {result['skipped_files']}", file=sys.stderr)
    print(f"  Output: {args.output}", file=sys.stderr)


if __name__ == "__main__":
    main()
