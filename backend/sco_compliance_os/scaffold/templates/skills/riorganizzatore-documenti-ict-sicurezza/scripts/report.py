#!/usr/bin/env python3
"""Phase 5: Reporting - Generate gap analysis, compliance maps, executive summary."""

import argparse
import json
import os
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

try:
    import openpyxl  # type: ignore[import-untyped]
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

SCRIPT_DIR = Path(__file__).parent.parent / "references"


def load_normative_map(path=None):
    if path is None:
        path = SCRIPT_DIR / "normative_mapping.json"
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


# ─── Styling helpers ─────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1B3A5C")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
ALT_FILL = PatternFill("solid", fgColor="F2F6FA")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"))
BODY_FONT = Font(name="Arial", size=10)


def style_header_row(ws, row, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN_BORDER


def style_body_cell(c, row_idx=0):
    c.font = BODY_FONT
    c.border = THIN_BORDER
    c.alignment = Alignment(vertical="top", wrap_text=True)
    if row_idx % 2 == 1:
        c.fill = ALT_FILL


def status_fill(status):
    return {"MAPPED": GREEN_FILL, "GAP": RED_FILL, "PARTIAL": YELLOW_FILL,
            "MULTI-MATCH": YELLOW_FILL, "ORPHAN": RED_FILL}.get(status, None)


# ─── Report: Gap Analysis ────────────────────────────────────
def report_gap_analysis(mapping, output_dir):
    if not HAS_OPENPYXL:
        print("openpyxl required for XLSX reports", file=sys.stderr)
        return _gap_analysis_json(mapping, output_dir)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gap Analysis"

    headers = ["COD.DOC", "TITOLO", "PROCESSO", "TIPOLOGIA", "PRIORITÀ", "FNCS",
               "STATO", "FILE MAPPATI", "CONFIDENCE", "REQUISITI"]
    widths = [12, 40, 30, 12, 10, 8, 14, 40, 12, 45]

    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(headers))

    for ri, doc in enumerate(mapping["document_mapping"]):
        row = ri + 2
        vals = [
            doc["cod_doc"], doc["titolo"], doc["processo"], doc["tipologia"],
            doc["priorita"], doc["funzione_fncs"], doc["status"],
            "; ".join(f["file_name"] for f in doc["mapped_files"]) if doc["mapped_files"] else "—",
            max((f["confidence"] for f in doc["mapped_files"]), default=0),
            doc.get("requisiti_fonti", "")
        ]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=ci, value=v)
            style_body_cell(c, ri)
            sf = status_fill(doc["status"])
            if ci == 7 and sf:
                c.fill = sf

    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Summary sheet
    ws2 = wb.create_sheet("Riepilogo")
    summary_data = [
        ["RIEPILOGO GAP ANALYSIS", ""],
        ["Data analisi", mapping["mapping_date"]],
        ["Documenti totali nell'indice", mapping["summary"]["total_index_documents"]],
        ["File sorgente analizzati", mapping["summary"]["total_source_files"]],
        ["Copertura complessiva", f"{mapping['summary']['coverage_percent']}%"],
        ["", ""],
        ["STATO", "CONTEGGIO"],
    ]
    for status, count in mapping["summary"]["status_counts"].items():
        summary_data.append([status, count])
    summary_data.extend([
        ["Orfani", mapping["summary"]["orphan_count"]],
        ["Gruppi duplicati", mapping["summary"]["duplicate_groups"]],
        ["", ""],
        ["COPERTURA PER PRIORITÀ", ""],
        ["Priorità", "Copertura"],
    ])
    for prio, data in mapping.get("priority_coverage", {}).items():
        summary_data.append([prio, f"{data['covered']}/{data['total']} ({data['percent']}%)"])

    for ri, row_data in enumerate(summary_data, 1):
        for ci, val in enumerate(row_data, 1):
            c = ws2.cell(row=ri, column=ci, value=val)
            c.font = BODY_FONT
            if ri == 1:
                c.font = Font(bold=True, name="Arial", size=14, color="1B3A5C")

    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 30

    path = Path(output_dir) / "gap_analysis.xlsx"
    wb.save(path)
    print(f"  [✓] Gap Analysis: {path}")
    return path


def _gap_analysis_json(mapping, output_dir):
    path = Path(output_dir) / "gap_analysis.json"
    gaps = [d for d in mapping["document_mapping"] if d["status"] == "GAP"]
    with open(path, "w", encoding="utf-8") as f:
        json.dump({"gaps": gaps, "total_gaps": len(gaps), "summary": mapping["summary"]}, f, ensure_ascii=False, indent=2)
    print(f"  [✓] Gap Analysis (JSON): {path}")
    return path


# ─── Report: Compliance Map ──────────────────────────────────
def report_compliance_map(mapping, output_dir):
    if not HAS_OPENPYXL:
        print("openpyxl required", file=sys.stderr)
        return None

    norm_map = load_normative_map()
    doc_status = {d["cod_doc"]: d["status"] for d in mapping["document_mapping"]}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compliance Map"

    headers = ["FONTE", "RIFERIMENTO", "AREA/REQUISITO", "DOCUMENTI RICHIESTI", "STATO COPERTURA", "DOCUMENTI PRESENTI", "DOCUMENTI MANCANTI"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(headers))

    for ri, req in enumerate(norm_map):
        row = ri + 2
        doc_codes = [c.strip() for c in req["documenti_correlati"].split(",") if c.strip()]
        present = [c for c in doc_codes if doc_status.get(c, "GAP") in ("MAPPED", "MULTI-MATCH", "PARTIAL")]
        missing = [c for c in doc_codes if doc_status.get(c, "GAP") == "GAP"]
        coverage = "COMPLIANT" if not missing else ("PARTIAL" if present else "NON-COMPLIANT")

        vals = [
            req["fonte"], req["riferimento"], req["area_requisito"],
            ", ".join(doc_codes), coverage,
            ", ".join(present) if present else "—",
            ", ".join(missing) if missing else "—"
        ]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=ci, value=v)
            style_body_cell(c, ri)
            if ci == 5:
                c.fill = {"COMPLIANT": GREEN_FILL, "PARTIAL": YELLOW_FILL, "NON-COMPLIANT": RED_FILL}.get(coverage, None)

    widths = [12, 18, 35, 35, 16, 30, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    path = Path(output_dir) / "compliance_map.xlsx"
    wb.save(path)
    print(f"  [✓] Compliance Map: {path}")
    return path


# ─── Report: Inventory ───────────────────────────────────────
def report_inventory(mapping, classification_file, output_dir):
    if not HAS_OPENPYXL:
        return None

    try:
        with open(classification_file, "r", encoding="utf-8") as f:
            clf = json.load(f)
    except FileNotFoundError:
        clf = {"classifications": []}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario Documentale"

    headers = ["FILE", "COD.DOC", "TITOLO INDICE", "PROCESSO", "TIPOLOGIA", "FNCS",
               "PRIORITÀ", "CONFIDENCE", "LIVELLO", "NOTE REFS"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(headers))

    for ri, cl in enumerate(clf.get("classifications", [])):
        row = ri + 2
        c = cl["classification"]
        vals = [
            cl["file_name"], c.get("cod_doc", "—"), c.get("titolo", "—"),
            c.get("processo", "—"), c.get("tipologia", "—"), c.get("funzione_fncs", "—"),
            c.get("priorita", "—"), cl["confidence"], cl["confidence_level"],
            "; ".join(r["reference"] for r in cl.get("normative_refs_found", []))
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=ci, value=v)
            style_body_cell(cell, ri)

    widths = [35, 12, 40, 30, 12, 8, 10, 10, 14, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    path = Path(output_dir) / "inventory.xlsx"
    wb.save(path)
    print(f"  [✓] Inventory: {path}")
    return path


# ─── Report: Completion Map ──────────────────────────────────
def report_completion_map(mapping, output_dir):
    if not HAS_OPENPYXL:
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mappa Completamento"

    types = ["POL", "PRC", "PRO", "LG", "IO", "MAN", "PLB", "PIA", "REG", "REP", "DOC", "TOTALE"]
    headers = ["PROCESSO"] + types
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(headers))

    # Group by process
    proc_type_counts = defaultdict(lambda: defaultdict(lambda: {"total": 0, "covered": 0}))
    for doc in mapping["document_mapping"]:
        proc = doc["processo"]
        tip = doc["tipologia"]
        proc_type_counts[proc][tip]["total"] += 1
        if doc["status"] in ("MAPPED", "MULTI-MATCH", "PARTIAL"):
            proc_type_counts[proc][tip]["covered"] += 1

    processes = sorted(proc_type_counts.keys())
    for ri, proc in enumerate(processes):
        row = ri + 2
        ws.cell(row=row, column=1, value=proc).font = Font(bold=True, name="Arial", size=10)
        total_all, covered_all = 0, 0
        for ti, tip in enumerate(types[:-1]):
            data = proc_type_counts[proc][tip]
            total_all += data["total"]
            covered_all += data["covered"]
            c = ws.cell(row=row, column=ti + 2)
            if data["total"] > 0:
                c.value = f"{data['covered']}/{data['total']}"
                c.fill = GREEN_FILL if data["covered"] == data["total"] else (YELLOW_FILL if data["covered"] > 0 else RED_FILL)
            else:
                c.value = "—"
            style_body_cell(c, ri)
        # Total column
        c = ws.cell(row=row, column=len(types) + 1, value=f"{covered_all}/{total_all}")
        c.font = Font(bold=True, name="Arial", size=10)
        c.border = THIN_BORDER

    ws.column_dimensions["A"].width = 35
    for i in range(2, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 10

    path = Path(output_dir) / "completion_map.xlsx"
    wb.save(path)
    print(f"  [✓] Completion Map: {path}")
    return path


# ─── Report: Executive Summary (JSON) ────────────────────────
def report_executive_summary(mapping, output_dir):
    s = mapping["summary"]
    gaps_alta = [d for d in mapping["document_mapping"] if d["status"] == "GAP" and d["priorita"] == "ALTA"]

    summary = {
        "title": "Executive Summary - Stato Documentazione Sicurezza ICT",
        "date": mapping["mapping_date"],
        "overall_coverage": f"{s['coverage_percent']}%",
        "total_documents_required": s["total_index_documents"],
        "documents_present": s["status_counts"].get("MAPPED", 0) + s["status_counts"].get("MULTI-MATCH", 0),
        "documents_partial": s["status_counts"].get("PARTIAL", 0),
        "documents_missing": s["status_counts"].get("GAP", 0),
        "orphan_files": s["orphan_count"],
        "priority_coverage": mapping.get("priority_coverage", {}),
        "critical_gaps": [{"cod_doc": g["cod_doc"], "titolo": g["titolo"], "processo": g["processo"],
                          "requisiti": g.get("requisiti_fonti", "")} for g in gaps_alta[:20]],
        "process_coverage": mapping.get("process_coverage", {}),
        "recommendations": _generate_recommendations(mapping)
    }

    path = Path(output_dir) / "executive_summary.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    print(f"  [✓] Executive Summary: {path}")
    return path


def _generate_recommendations(mapping):
    recs = []
    alta = mapping.get("priority_coverage", {}).get("ALTA", {})
    if alta.get("percent", 100) < 80:
        recs.append(f"CRITICO: Copertura documenti ALTA priorità solo al {alta['percent']}%. Prioritizzare la redazione dei {alta['total'] - alta['covered']} documenti mancanti.")

    gaps_by_process = defaultdict(int)
    for d in mapping["document_mapping"]:
        if d["status"] == "GAP" and d["priorita"] == "ALTA":
            gaps_by_process[d["processo"]] += 1
    worst = sorted(gaps_by_process.items(), key=lambda x: x[1], reverse=True)[:3]
    for proc, count in worst:
        recs.append(f"Area '{proc}': {count} documenti obbligatori mancanti.")

    if mapping["summary"]["orphan_count"] > 0:
        recs.append(f"{mapping['summary']['orphan_count']} file non classificabili richiedono revisione manuale.")
    if mapping["summary"]["duplicate_groups"] > 0:
        recs.append(f"{mapping['summary']['duplicate_groups']} gruppi di file duplicati da consolidare.")

    return recs


# ─── Main ────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Phase 5: Generate reports")
    parser.add_argument("mapping", help="Path to mapping.json")
    parser.add_argument("--type", "-t", default="all", choices=["gap", "compliance", "inventory", "completion", "summary", "all"])
    parser.add_argument("--classification", default="classification.json", help="Path to classification.json (for inventory report)")
    parser.add_argument("--output", "-o", default="reports", help="Output directory")
    args = parser.parse_args()

    with open(args.mapping, "r", encoding="utf-8") as f:
        mapping = json.load(f)

    os.makedirs(args.output, exist_ok=True)
    print(f"\nGenerating reports in: {args.output}/\n")

    report_type = args.type

    if report_type in ("gap", "all"):
        report_gap_analysis(mapping, args.output)
    if report_type in ("compliance", "all"):
        report_compliance_map(mapping, args.output)
    if report_type in ("inventory", "all"):
        report_inventory(mapping, args.classification, args.output)
    if report_type in ("completion", "all"):
        report_completion_map(mapping, args.output)
    if report_type in ("summary", "all"):
        report_executive_summary(mapping, args.output)

    print(f"\nReporting complete.")


if __name__ == "__main__":
    main()
